import os
import re
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
from telegram import Update
from telegram.ext import ApplicationBuilder, CommandHandler, ContextTypes

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

BOT_TOKEN = os.getenv("BOT_TOKEN")
GROUP_CHAT_ID = int(os.getenv("GROUP_CHAT_ID"))
EXCEL_PATH = BASE_DIR / "PiedraRoseta.xlsx"

PENDING_UPDATES = {}

PERSON_TO_COL = {
    "carlitos": "B",
    "yago": "C",
    "nikola": "D",
}

MONTH_ALIASES = {
    "enero": "Enero",
    "febrero": "Febrero",
    "marzo": "Marzo",
    "abril": "Abril",
    "mayo": "Mayo",
    "junio": "Junio",
    "julio": "Julio",
    "agosto": "Agosto",
    "septiembre": "Septiembre",
    "setiembre": "Septiembre",
    "octubre": "Octubre",
    "noviembre": "Noviembre",
    "diciembre": "Diciembre",
}


def normalize_text(text: str) -> str:
    text = text.strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text

def prepare_update(month: str, person: str, exercise: str, new_weight: float):
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"N existe el archivo {EXCEL_PATH.name}")

    person_key = normalize_text(person)
    if person_key not in PERSON_TO_COL:
        raise ValueError(f"Persona no válida: {person}")

    col = PERSON_TO_COL[person_key]

    wb = load_workbook(EXCEL_PATH)
    ws = get_sheet_by_month(wb, month)
    row = find_exercise_row(ws, exercise)

    cell = f"{col}{row}"
    old_value = ws[cell].value

    old_numeric = None
    if isinstance(old_value, (int, float)):
        old_numeric = float(old_value)
    else:
        try:
            if old_value is not None and str(old_value).strip() != "":
                old_numeric = float(str(old_value).replace(",", "."))
        except ValueError:
            old_numeric = None

    suspicious_increase = False
    if old_numeric is not None and old_numeric > 0:
        suspicious_increase = new_weight > old_numeric * 1.20

    return {
        "sheet": ws.title,
        "row": row,
        "cell": cell,
        "old_value": old_value,
        "old_numeric": old_numeric,
        "new_value": new_weight,
        "month": month,
        "person": person,
        "exercise": exercise,
        "col": col,
        "suspicious_increase": suspicious_increase,
    }

def parse_weight(value: str) -> float:
    return float(value.replace(",", "."))

def apply_update(prepared_update: dict):
    wb = load_workbook(EXCEL_PATH)
    ws = wb[prepared_update["sheet"]]

    cell = prepared_update["cell"]
    ws[cell] = prepared_update["new_value"]

    wb.save(EXCEL_PATH)

    return {
        "sheet": prepared_update["sheet"],
        "cell": cell,
        "old_value": prepared_update["old_value"],
        "new_value": prepared_update["new_value"],
    }

def get_sheet_by_month(workbook, month_input: str):
    month_key = normalize_text(month_input)
    if month_key not in MONTH_ALIASES:
        raise ValueError(f"Mes no reconocido: {month_input}")

    sheet_name = MONTH_ALIASES[month_key]
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"No existe la hoja '{sheet_name}' en el Excel")

    return workbook[sheet_name]


def find_exercise_row(ws, exercise_input: str) -> int:
    target = normalize_text(exercise_input)

    for row in range(1, ws.max_row + 1):
        value = ws[f"A{row}"].value
        if value is None:
            continue

        current = normalize_text(str(value))
        if current == target:
            return row

    raise ValueError(f"No encuentro el ejercicio '{exercise_input}' en la hoja '{ws.title}'")


def update_excel(month: str, person: str, exercise: str, new_weight: float):
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"No existe el archivo {EXCEL_PATH.name}")

    person_key = normalize_text(person)
    if person_key not in PERSON_TO_COL:
        raise ValueError(f"Persona no válida: {person}")

    col = PERSON_TO_COL[person_key]

    wb = load_workbook(EXCEL_PATH)
    ws = get_sheet_by_month(wb, month)
    row = find_exercise_row(ws, exercise)

    cell = f"{col}{row}"
    old_value = ws[cell].value

    old_numeric = None
    if isinstance(old_value, (int, float)):
        old_numeric = float(old_value)
    else:
        try:
            if old_value is not None and str(old_value).strip() != "":
                old_numeric = float(str(old_value).replace(",", "."))
        except ValueError:
            old_numeric = None

    suspicious_increase = False
    if old_numeric is not None and old_numeric > 0:
        suspicious_increase = new_weight > old_numeric * 1.20

    ws[cell] = new_weight
    wb.save(EXCEL_PATH)

    return {
        "sheet": ws.title,
        "cell": cell,
        "old_value": old_value,
        "new_value": new_weight,
        "suspicious_increase": suspicious_increase,
    }

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "Tranquilo no soy una notificación de Sofía Gonzalez solo que estoy funcionando :) \n\n"
        "Pruea esto:\n"
        "/hito marzo carlitos 72.5 press banca"
    )

async def chatid(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(f"Chat ID: {update.effective_chat.id}")

async def hito(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        if len(context.args) < 4:
            raise ValueError("Formato incorrecto")

        month = context.args[0]
        person = context.args[1]
        weight = parse_weight(context.args[2])
        exercise = " ".join(context.args[3:]).strip()

        prepared = prepare_update(month, person, exercise, weight)

        user_id = update.effective_user.id

        if prepared["suspicious_increase"]:
            PENDING_UPDATES[user_id] = prepared

            old_value = prepared["old_value"]
            await update.message.reply_text(
                f"⚠️ Atención: el nuevo peso ({weight} kg) supera en más de un 20% "
                f"el valor anterior ({old_value}).\n\n"
                f"No he guardado el cambio todavía.\n"
                f"Escribe /confirmar para guardarlo o /cancelar para anularlo."
            )
            return

        result = apply_update(prepared)

        await update.message.reply_text(
            f"✅ Actualizado correctamente\n"
            f"Hoja: {result['sheet']}\n"
            f"Celda: {result['cell']}\n"
            f"Antes: {result['old_value']}\n"
            f"Ahora: {result['new_value']} kg"
        )

        await context.bot.send_message(
            chat_id=GROUP_CHAT_ID,
            text=(
                f"🏆 Nuevo hito en el gimnasio\n"
                f"{person.title()} ha registrado {weight} kg en {exercise} ({result['sheet']})."
            ),
        )

    except Exception as e:
        print("ERROR EN /hito:", e)
        await update.message.reply_text(
            "No pude actualizar el Excel.\n"
            f"Error: {e}\n\n"
            "Formato esperado:\n"
            "/hito marzo carlitos 72.5 press banca"
        )

async def confirmar(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id

    if user_id not in PENDING_UPDATES:
        await update.message.reply_text("No tienes ninguna actualización pendiente de confirmar.")
        return

    prepared = PENDING_UPDATES.pop(user_id)
    result = apply_update(prepared)

    await update.message.reply_text(
        f"✅ Cambio confirmado y guardado\n"
        f"Hoja: {result['sheet']}\n"
        f"Celda: {result['cell']}\n"
        f"Antes: {result['old_value']}\n"
        f"Ahora: {result['new_value']} kg"
    )

    await context.bot.send_message(
        chat_id=GROUP_CHAT_ID,
        text=(
            f"🏆 Nuevo hito en el gimnasio\n"
            f"{prepared['person'].title()} ha registrado {prepared['new_value']} kg "
            f"en {prepared['exercise']} ({result['sheet']}).\n"
            f"✅ Cambio confirmado manualmente."
        ),
    )

async def cancelar(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id

    if user_id in PENDING_UPDATES:
        PENDING_UPDATES.pop(user_id)
        await update.message.reply_text("❌ Actualización cancelada. No se ha modificado el Excel.")
    else:
        await update.message.reply_text("No tienes ninguna actualización pendiente.")

def main():
    if not BOT_TOKEN:
        raise RuntimeError("No se encontró BOT_TOKEN en el archivo .env")

    if not EXCEL_PATH.exists():
        raise RuntimeError(f"No se encontró el Excel: {EXCEL_PATH}")

    app = ApplicationBuilder().token(BOT_TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("hito", hito))
    app.add_handler(CommandHandler("chatid", chatid))
    app.add_handler(CommandHandler("confirmar", confirmar))
    app.add_handler(CommandHandler("cancelar", cancelar))


    print("Bot arrancado correctamente")
    app.run_polling()


if __name__ == "__main__":
    main()